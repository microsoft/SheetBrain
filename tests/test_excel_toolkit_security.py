"""Workbook-level security regression tests."""

import os
import tempfile
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

from core.agent import MAX_PREVIEW_CELLS, SheetBrain
from modules.execution import ALLOWED_EXCEL_HELPERS, HELPER_CALL_LIMITS
from utils.excel_toolkit import ExcelToolkit, MAX_MUTATION_COUNT, MAX_WORKSHEET_ROWS


class ExcelToolkitSecurityTests(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.input_path = os.path.join(self.temp_directory.name, "input.xlsx")
        self.workbook = Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "Data"
        self.workbook.save(self.input_path)
        self.toolkit = ExcelToolkit(self.workbook, self.input_path)

    def tearDown(self):
        self.temp_directory.cleanup()

    def test_model_helper_allowlist_is_fully_rate_limited(self):
        helper_names = set(self.toolkit.get_helper_functions_dict())

        self.assertEqual(ALLOWED_EXCEL_HELPERS, set(HELPER_CALL_LIMITS))
        self.assertTrue(ALLOWED_EXCEL_HELPERS <= helper_names)
        self.assertTrue({"get_sheet", "get_sheet_as_dataframe", "save_plot_to_excel"}.isdisjoint(
            ALLOWED_EXCEL_HELPERS
        ))

    def test_generic_writes_and_copy_neutralize_formulas(self):
        self.toolkit.set_cell_value("Data", "A1", "=WEBSERVICE(\"https://example.test\")")
        self.toolkit.set_range_values("Data", "A2", [["+cmd", "plain"]])
        self.toolkit.copy_range("Data", "A1:B2", "Data", "C1")

        self.assertEqual("'=WEBSERVICE(\"https://example.test\")", self.sheet["A1"].value)
        self.assertEqual("'+cmd", self.sheet["A2"].value)
        self.assertEqual("plain", self.sheet["B2"].value)
        self.assertEqual("'=WEBSERVICE(\"https://example.test\")", self.sheet["C1"].value)

    def test_explicit_formula_policy_is_enforced(self):
        self.toolkit.add_formula("Data", "A1", "SUM(B1:B2)")
        self.assertEqual("=SUM(B1:B2)", self.sheet["A1"].value)

        with self.assertRaises(Exception):
            self.toolkit.add_formula("Data", "A2", "=HYPERLINK(\"https://example.test\", \"x\")")

    def test_large_ranges_and_mutations_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "cell limit"):
            self.toolkit.inspector("A1:A100001", "Data")
        with self.assertRaises(Exception):
            self.toolkit.insert_rows("Data", 1, MAX_MUTATION_COUNT + 1)
        with self.assertRaises(Exception):
            self.toolkit.insert_rows("Data", MAX_WORKSHEET_ROWS + 1)
        with self.assertRaises(Exception):
            self.toolkit.set_cell_value("Data", "A1048577", "outside")

    def test_copy_range_checks_destination_before_writing(self):
        self.sheet["A1"] = "first"
        self.sheet["B1"] = "second"

        with self.assertRaises(Exception):
            self.toolkit.copy_range("Data", "A1:B1", "Data", "XFD1")

        self.assertIsNone(self.sheet["XFD1"].value)

    def test_sparse_oversized_sheet_search_is_rejected(self):
        self.sheet["XFD1048576"] = "needle"

        with self.assertRaisesRegex(ValueError, "Search range exceeds"):
            self.toolkit.search("needle", "Data")

    def test_save_uses_unique_non_overwriting_files(self):
        fixed_legacy_path = os.path.join(self.temp_directory.name, "input_output.xlsx")
        with open(fixed_legacy_path, "w", encoding="utf-8") as legacy_file:
            legacy_file.write("do not overwrite")

        first_path = self.toolkit.save_workbook()
        second_path = self.toolkit.save_workbook()

        self.assertNotEqual(first_path, second_path)
        self.assertTrue(os.path.exists(first_path))
        self.assertTrue(os.path.exists(second_path))
        with open(fixed_legacy_path, encoding="utf-8") as legacy_file:
            self.assertEqual("do not overwrite", legacy_file.read())
        saved_workbook = load_workbook(first_path, read_only=True)
        try:
            self.assertEqual(["Data"], saved_workbook.sheetnames)
        finally:
            saved_workbook.close()


class PreviewSecurityTests(unittest.TestCase):
    def test_accepts_normal_workbook_archive(self):
        with tempfile.TemporaryDirectory() as temp_directory:
            workbook_path = os.path.join(temp_directory, "normal.xlsx")
            Workbook().save(workbook_path)

            SheetBrain._validate_workbook_archive(workbook_path)

    def test_preview_respects_global_cell_limit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["XFD1048576"] = "far away"
        agent = SheetBrain.__new__(SheetBrain)

        result = agent._get_sheet_preview_with_token_limit(
            sheet,
            token_budget=1000000,
            max_rows=1000000,
            max_cols=1000000,
            max_cells=MAX_PREVIEW_CELLS,
        )

        self.assertLessEqual(result["cells_scanned"], MAX_PREVIEW_CELLS)

    def test_rejects_suspiciously_compressed_workbook_archive(self):
        with tempfile.TemporaryDirectory() as temp_directory:
            archive_path = os.path.join(temp_directory, "bomb.xlsx")
            with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("xl/worksheets/sheet1.xml", "0" * 1000000)

            with self.assertRaisesRegex(ValueError, "compression ratio"):
                SheetBrain._validate_workbook_archive(archive_path)

    def test_public_execution_result_removes_sensitive_diagnostics(self):
        result = SheetBrain._public_execution_result({
            "success": True,
            "answer": "approved answer",
            "total_turns": 1,
            "conversation_history": [{"content": "secret workbook data"}],
            "execution_summary": {
                "total_code_executions": 1,
                "execution_steps": [{"code": "secret code", "result": "secret output"}],
                "final_answer": "duplicate sensitive answer",
            },
        })

        self.assertEqual("approved answer", result["answer"])
        self.assertNotIn("conversation_history", result)
        self.assertNotIn("execution_steps", result["execution_summary"])
        self.assertNotIn("final_answer", result["execution_summary"])

    def test_public_validation_result_removes_sensitive_text(self):
        result = SheetBrain._public_validation_result({
            "validation_passed": False,
            "confidence_score": 0.5,
            "requires_reexecution": True,
            "issues_found": ["secret workbook value"],
            "improvement_feedback": "secret feedback",
            "final_assessment": "secret assessment",
        })

        self.assertEqual({
            "validation_passed": False,
            "confidence_score": 0.5,
            "requires_reexecution": True,
        }, result)

    def test_editable_workbook_preserves_formulas(self):
        with tempfile.TemporaryDirectory() as temp_directory:
            workbook_path = os.path.join(temp_directory, "formulas.xlsx")
            workbook = Workbook()
            workbook.active["A1"] = "=SUM(1,2)"
            workbook.save(workbook_path)

            agent = SheetBrain.__new__(SheetBrain)
            agent.excel_path = workbook_path
            agent.code_globals = {}
            agent._setup_excel_libraries()
            try:
                self.assertEqual("=SUM(1,2)", agent.workbook.active["A1"].value)
            finally:
                agent.preview_workbook.close()
                agent.workbook.close()


if __name__ == "__main__":
    unittest.main()